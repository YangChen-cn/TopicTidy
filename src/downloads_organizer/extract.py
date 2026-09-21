from __future__ import annotations

import collections
import plistlib
import re
import subprocess
from pathlib import Path
from urllib.parse import unquote, urlparse

from .config import SUPPORTED_TEXT_SUFFIXES
from .models import Extracted

WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9_-]{2,}|[\u4e00-\u9fff]{2,}")
STOPWORDS = {
    "the", "and", "for", "with", "from", "this", "that", "chapter", "introduction",
    "lecture", "notes", "slide", "slides", "document", "file", "course", "week",
    "的", "了", "和", "以及", "课程", "讲义", "章节", "介绍",
}


def source_urls(path: Path) -> list[str]:
    try:
        raw = subprocess.run(
            ["xattr", "-p", "com.apple.metadata:kMDItemWhereFroms", str(path)],
            check=True, capture_output=True, timeout=3,
        ).stdout
        value = plistlib.loads(raw)
        if isinstance(value, list):
            return [str(x) for x in value if str(x).startswith(("http://", "https://"))]
    except (subprocess.SubprocessError, OSError, plistlib.InvalidFileException):
        pass
    try:
        result = subprocess.run(
            ["mdls", "-raw", "-name", "kMDItemWhereFroms", str(path)],
            check=True, capture_output=True, text=True, timeout=3,
        ).stdout
        return re.findall(r'"(https?://[^"\\]+)', result)
    except (subprocess.SubprocessError, OSError):
        return []


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValueError("PDF 已加密") from exc
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs]
    blocks.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(blocks)


def _pptx(path: Path) -> str:
    from pptx import Presentation
    deck = Presentation(path)
    return "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes
        if hasattr(shape, "text") and shape.text
    )


def _plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def keywords(text: str, limit: int = 12) -> list[str]:
    tokens = [x.lower() for x in WORD_RE.findall(text)]
    counts = collections.Counter(x for x in tokens if x not in STOPWORDS and not x.isdigit())
    return [word for word, _ in counts.most_common(limit)]


def extract(path: Path, max_chars: int) -> Extracted:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_TEXT_SUFFIXES:
        return Extracted()
    try:
        if suffix == ".pdf":
            text = _pdf(path)
        elif suffix == ".docx":
            text = _docx(path)
        elif suffix == ".pptx":
            text = _pptx(path)
        else:
            text = _plain(path)
        text = re.sub(r"[ \t]+", " ", text).strip()
        truncated = len(text) > max_chars
        if truncated:
            half = max_chars // 2
            text = text[:half] + "\n…\n" + text[-half:]
        lines = [line.strip(" #\t") for line in text.splitlines() if len(line.strip()) >= 3]
        title = lines[0][:180] if lines else ""
        keys = keywords(text)
        representative = " ".join(lines[:5])[:600]
        return Extracted(text=text, title=title, keywords=keys, summary=representative, truncated=truncated)
    except Exception as exc:
        return Extracted(error=f"{type(exc).__name__}: {exc}")


def url_tokens(urls: list[str]) -> set[str]:
    values: set[str] = set()
    for url in urls:
        parsed = urlparse(url)
        material = unquote(parsed.path + " " + parsed.query)
        values.update(x.lower() for x in WORD_RE.findall(material) if x.lower() not in STOPWORDS)
    return values

