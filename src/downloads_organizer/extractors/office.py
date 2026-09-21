from __future__ import annotations

from pathlib import Path

from .base import ExtractionContext, finalize_text
from ..models import Extracted


def _append_with_budget(parts: list[str], value: str, remaining: int) -> tuple[int, bool]:
    if remaining <= 0:
        return 0, bool(value)
    value = value.strip()
    if not value:
        return remaining, False
    selected = value[:remaining]
    parts.append(selected)
    return remaining - len(selected), len(value) > len(selected)


class DocxExtractor:
    name = "docx"
    version = "1"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def extract(self, path: Path, context: ExtractionContext) -> Extracted:
        from docx import Document

        document = Document(path)
        parts: list[str] = []
        remaining, truncated = context.max_chars, False
        for paragraph in document.paragraphs:
            remaining, cut = _append_with_budget(parts, paragraph.text, remaining)
            truncated |= cut
            if remaining <= 0:
                break
        if remaining > 0:
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        remaining, cut = _append_with_budget(parts, cell.text, remaining)
                        truncated |= cut
                        if remaining <= 0:
                            break
                    if remaining <= 0:
                        break
                if remaining <= 0:
                    break
        combined = "\n".join(parts)
        return finalize_text(combined[:context.max_chars], truncated=truncated or len(combined) > context.max_chars)


class PptxExtractor:
    name = "pptx"
    version = "1"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    def extract(self, path: Path, context: ExtractionContext) -> Extracted:
        from pptx import Presentation

        deck = Presentation(path)
        parts: list[str] = []
        remaining, truncated = context.max_chars, False
        for slide in deck.slides:
            for shape in slide.shapes:
                if not hasattr(shape, "text") or not shape.text:
                    continue
                remaining, cut = _append_with_budget(parts, shape.text, remaining)
                truncated |= cut
                if remaining <= 0:
                    break
            if remaining <= 0:
                break
        combined = "\n".join(parts)
        return finalize_text(combined[:context.max_chars], truncated=truncated or len(combined) > context.max_chars)
