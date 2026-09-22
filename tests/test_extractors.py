from __future__ import annotations

from pathlib import Path

from docx import Document

from downloads_organizer.extractors import ExtractionContext, ExtractorRegistry
from downloads_organizer.extractors.office import DocxExtractor
from downloads_organizer.extractors.pdf import PdfExtractor, sampled_page_indices
from downloads_organizer.models import Extracted
from downloads_organizer.scanner import scan
from downloads_organizer.text_features import TEXT_FEATURE_VERSION

from conftest import put


class CustomExtractor:
    name = "custom"
    version = "7"

    def supports(self, path: Path) -> bool:
        return path.suffix == ".custom"

    def extract(self, path: Path, context: ExtractionContext) -> Extracted:
        return Extracted(text="plugged in", title="Custom Title", keywords=["plugged"])


def test_scanner_accepts_pluggable_extractor_without_format_library_dependency(workspace):
    settings, db = workspace
    put(settings.downloads, "sample.custom", "opaque")
    registry = ExtractorRegistry([CustomExtractor()])

    stats = scan(db, settings, extractors=registry)

    assert stats["scanned"] == 1
    feature = db.conn.execute("SELECT extractor_version,text,title FROM features").fetchone()
    assert dict(feature) == {
        "extractor_version": f"custom:7;text-features:{TEXT_FEATURE_VERSION}",
        "text": "plugged in",
        "title": "Custom Title",
    }


def test_docx_extractor_reads_real_document_with_budget(tmp_path):
    path = tmp_path / "Grid Storage Design v2.docx"
    document = Document()
    document.add_heading("Grid Storage Design", 0)
    document.add_paragraph("battery storage " * 200)
    document.save(path)

    result = DocxExtractor().extract(path, ExtractionContext(max_chars=180))

    assert result.title == "Grid Storage Design"
    assert len(result.text) <= 180
    assert result.truncated


class FakePage:
    def __init__(self, index: int, accessed: list[int]):
        self.index = index
        self.accessed = accessed

    def extract_text(self) -> str:
        self.accessed.append(self.index)
        return f"PAGE-{self.index} " + ("representative text " * 100)


class FakeReader:
    is_encrypted = False

    def __init__(self, count: int, accessed: list[int]):
        self.pages = [FakePage(index, accessed) for index in range(count)]


def test_large_pdf_samples_front_middle_and_tail_without_reading_every_page(tmp_path):
    accessed: list[int] = []
    extractor = PdfExtractor(lambda _: FakeReader(100, accessed))

    result = extractor.extract(tmp_path / "large.pdf", ExtractionContext(max_chars=320))

    assert accessed == sampled_page_indices(100)
    assert accessed == [0, 1, 2, 25, 50, 74, 98, 99]
    assert "PAGE-0" in result.text
    assert "PAGE-50" in result.text
    assert "PAGE-99" in result.text
    assert len(result.text) <= 320
    assert result.truncated
